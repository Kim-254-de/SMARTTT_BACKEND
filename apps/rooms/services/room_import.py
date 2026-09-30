"""
Import room capacities from an admin-uploaded CSV/XLSX file.

Column names are matched loosely (case, spaces and punctuation ignored) so the
registry's own spreadsheet can be uploaded without reformatting. Only a room
code and a capacity column are required; rooms are matched by code
(case-insensitive) and created if they don't exist yet.
"""
import csv
import io
import re
from typing import Any, Dict, List, Optional

from django.db import transaction

from apps.rooms.models import Room

# Normalised header -> Room field. Normalisation lowercases and strips
# everything that isn't a letter or digit, so "Room Code", "room_code" and
# "ROOM-CODE" all become "roomcode".
HEADER_ALIASES = {
    "code": "code",
    "roomcode": "code",
    "room": "code",
    "roomno": "code",
    "roomnumber": "code",
    "venue": "code",
    "venuecode": "code",
    "hall": "code",
    "capacity": "capacity",
    "roomcapacity": "capacity",
    "venuecapacity": "capacity",
    "seats": "capacity",
    "seatingcapacity": "capacity",
    "noofseats": "capacity",
    "numberofseats": "capacity",
    "maxstudents": "capacity",
    "size": "capacity",
    "name": "name",
    "roomname": "name",
    "venuename": "name",
    "building": "building",
    "block": "building",
    "location": "building",
    "floor": "floor",
    "level": "floor",
    "type": "room_type",
    "roomtype": "room_type",
    "venuetype": "room_type",
    "status": "status",
}

REQUIRED_FIELDS = {"code", "capacity"}

# How far down the sheet to look for the header row (files often start with a
# title or department name).
HEADER_SEARCH_ROWS = 15

MAX_CODE_LENGTH = Room._meta.get_field("code").max_length


class RoomImportError(Exception):
    """The file as a whole could not be read (bad format, no header row)."""


def _normalise(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _cell(value: Any) -> str:
    if value is None:
        return ""
    # Excel stores whole numbers as floats: 120 -> 120.0
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_rows(file_name: str, content: bytes) -> List[List[str]]:
    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""

    if ext == "csv":
        try:
            # utf-8-sig: Excel-saved CSVs start with a BOM that would hide the header.
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        return [[_cell(c) for c in row] for row in csv.reader(io.StringIO(text))]

    if ext == "xlsx":
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise RoomImportError("Could not open the Excel file. Save it as .xlsx or .csv and try again.") from exc
        # Use the first sheet that has a recognisable header row.
        for sheet in workbook.worksheets:
            rows = [[_cell(c) for c in row] for row in sheet.iter_rows(values_only=True)]
            if _find_header(rows) is not None:
                return rows
        return []

    raise RoomImportError("Unsupported file type. Upload a .csv or .xlsx file.")


def _find_header(rows: List[List[str]]) -> Optional[tuple]:
    """Return (row_index, {column_index: field}) for the first header row found."""
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        columns = {}
        for col, value in enumerate(row):
            field = HEADER_ALIASES.get(_normalise(value))
            if field and field not in columns.values():
                columns[col] = field
        if REQUIRED_FIELDS <= set(columns.values()):
            return index, columns
    return None


def _match_choice(value: str, choices) -> Optional[str]:
    """Map free text like 'Lecture Hall', 'LAB' or 'maintenance' onto a choice value."""
    key = _normalise(value)
    if not key:
        return None
    for choice_value, label in choices:
        if key in (_normalise(choice_value), _normalise(label)):
            return choice_value
    for choice_value, label in choices:
        if _normalise(label).startswith(key) or key.startswith(_normalise(choice_value)):
            return choice_value
    return None


def _parse_capacity(raw: str) -> int:
    try:
        capacity = int(float(raw.replace(",", "")))
    except ValueError:
        raise ValueError(f"Capacity '{raw}' is not a number.")
    if capacity <= 0:
        raise ValueError("Capacity must be greater than 0.")
    return capacity


def import_rooms(file_name: str, content: bytes) -> Dict[str, Any]:
    rows = _read_rows(file_name, content)
    header = _find_header(rows)
    if header is None:
        raise RoomImportError(
            "Could not find a header row. The file needs at least a room code column "
            "(e.g. 'Room', 'Code', 'Venue') and a capacity column (e.g. 'Capacity', 'Seats')."
        )
    header_index, columns = header

    existing = {room.code.lower(): room for room in Room.objects.all()}
    created, updated, unchanged = [], [], []
    errors = []
    seen_codes = {}

    with transaction.atomic():
        for offset, row in enumerate(rows[header_index + 1:], start=header_index + 2):
            values = {field: row[col] if col < len(row) else "" for col, field in columns.items()}
            if not any(values.values()):
                continue  # blank line

            code = values["code"]
            if not code:
                errors.append({"row": offset, "code": "", "error": "Room code is missing."})
                continue
            if len(code) > MAX_CODE_LENGTH:
                errors.append({"row": offset, "code": code, "error": f"Room code is longer than {MAX_CODE_LENGTH} characters."})
                continue
            if code.lower() in seen_codes:
                errors.append({
                    "row": offset,
                    "code": code,
                    "error": f"Duplicate of row {seen_codes[code.lower()]}; this row was skipped.",
                })
                continue
            seen_codes[code.lower()] = offset

            try:
                capacity = _parse_capacity(values["capacity"])
            except ValueError as exc:
                errors.append({"row": offset, "code": code, "error": str(exc)})
                continue

            changes = {"capacity": capacity, "capacity_confirmed": True}
            for field in ("name", "building", "floor"):
                if values.get(field):
                    changes[field] = values[field][: Room._meta.get_field(field).max_length]
            if values.get("room_type"):
                room_type = _match_choice(values["room_type"], Room.Type.choices)
                if room_type is None:
                    errors.append({
                        "row": offset,
                        "code": code,
                        "error": f"Unknown room type '{values['room_type']}' ignored; capacity was still saved.",
                    })
                else:
                    changes["room_type"] = room_type
            if values.get("status"):
                status = _match_choice(values["status"], Room.Status.choices)
                if status is None:
                    errors.append({
                        "row": offset,
                        "code": code,
                        "error": f"Unknown status '{values['status']}' ignored; capacity was still saved.",
                    })
                else:
                    changes["status"] = status

            room = existing.get(code.lower())
            if room is None:
                changes.setdefault("name", code)
                room = Room.objects.create(code=code, **changes)
                existing[code.lower()] = room
                created.append(code)
                continue

            dirty = [field for field, value in changes.items() if getattr(room, field) != value]
            if not dirty:
                unchanged.append(room.code)
                continue
            for field in dirty:
                setattr(room, field, changes[field])
            room.save(update_fields=dirty + ["updated_at"])
            updated.append(room.code)

    return {
        "created": len(created),
        "updated": len(updated),
        "unchanged": len(unchanged),
        "errors": errors,
        "created_codes": created,
        "updated_codes": updated,
    }


def capacity_summary() -> Dict[str, Any]:
    """Rooms still running on a placeholder capacity, for the admin dashboard."""
    unconfirmed = list(
        Room.objects.filter(capacity_confirmed=False).exclude(code="TBA").order_by("code").values_list("code", flat=True)
    )
    total = Room.objects.exclude(code="TBA").count()
    return {
        "total": total,
        "confirmed": total - len(unconfirmed),
        "unconfirmed": len(unconfirmed),
        "unconfirmed_codes": unconfirmed,
    }
