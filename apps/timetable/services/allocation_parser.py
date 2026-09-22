"""
Parser for departmental course-allocation documents (.docx / .doc / .pdf).

Each allocation document is a sequence of programme sections, e.g.

    BACHELOR OF EDUCATION (ARTS): Y1S1 ...            <- heading (paragraph or merged row)
    Group A: ENGL/LIT                                  <- group -> subject combinations
    Group B: KISW/RELI, PE/RELI
    CODE | TITLE | L | P | CF | LECTURER | CONTACT     <- header row
    COSC 103 Group A | Intro to ICT | ... | Joseph Mutwiri (FT) | 07...
    COSC 103 Group B | Intro to ICT | ... | Michael Mutisya (FT) | 07...

The same unit is routinely allocated to different lecturers per programme,
year or group, so every row carries the section context it was found in
(year/semester, group letter and that group's subject combinations). The
matcher (allocation_matcher.py) uses that context to pin each row to the
exact master-timetable slots it belongs to.
"""
from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

# "COSC 103", "COSC 00104", "MATH100", "HIST  223" followed by optional group noise.
_CODE_RE = re.compile(r"^\s*([A-Z]{3,4})\s*(\d{2,5})(?!\d)(.*)$", re.DOTALL)
# Group suffix on a code cell: "G A", "GA", "Group A", "(GRP K)", "/Group V/A&B", "GR.M", "G D2".
_GROUP_RE = re.compile(r"(?<![A-Z])(?:GROUP|GRP|GR\.?|G)\s*([A-Z]\d?)(?![A-Z])")
# Section headings: "Y1S1", "Y2 S2", "YEAR 1 SEMESTER 2".
_YEAR_SEM_RE = re.compile(r"\bY(?:EAR)?\s*(\d)\s*,?\s*S(?:EM(?:ESTER)?)?\s*(\d)\b", re.IGNORECASE)
# "Group A: ENGL/LIT" / "Group D: GEO/PE (0745...-collins), KISW/GEO (...)" definitions.
_GROUP_DEF_RE = re.compile(r"^\s*GROUP\s+([A-Z]\d?)\s*[:\-–]\s*(.+)$", re.IGNORECASE)
_PHONE_RE = re.compile(r"(?:\+?254|0)[17]\d{8}")
# Department service placeholders written where a lecturer name is not yet known:
# "DCOMP", "DEDU", "DHUM- 1.2", "DDANR", "TBA".
_PLACEHOLDER_RE = re.compile(
    r"^(?:D[A-Z]{2,5}\b[\s\-\d\.]*|TBA|TBD|N/?A|-+|HUMANITIES|SOCIAL\s+SCIENCES?|CO-?ORDINATION|MOVED\b.*|DEPARTMENT\b.*|DEPT\b.*)$",
    re.IGNORECASE,
)
_PROGRAMME_RE = re.compile(r"^\s*(BACHELOR|DIPLOMA|CERTIFICATE|MASTER|DOCTOR|POSTGRADUATE|PHD)\b", re.IGNORECASE)
_ODEL_RE = re.compile(r"(?<![A-Z])ODEL\b", re.IGNORECASE)
# "COURSE ALLOCATION SEP-DEC 2026 (ODEL)" / "(REGULAR)" part titles within one document.
_PART_RE = re.compile(r"^\s*COURSE\s+ALLOCATION\b.*\((ODEL|REGULAR)\)", re.IGNORECASE)
_DEPARTMENT_RE = re.compile(r"^\s*(DEPARTMENT\s+OF\s+[A-Z&,\s]+?)\s*$", re.IGNORECASE)


def _unique_cells(row):
    """Filters out python-docx duplicate references to merged cells."""
    unique = []
    for cell in row.cells:
        if not unique or cell._tc is not unique[-1]._tc:
            unique.append(cell)
    return unique


def parse_code_cell(text: str) -> tuple[str, str, list[str]] | None:
    """
    'COSC 103 Group A' -> ('COSC103', 'A', []);  'MATH 122 (GRP K)' -> ('MATH122', 'K', [])
    'BOTA 101 G A' -> ('BOTA101', 'A', []);      'COSC 00104' -> ('COSC00104', '', [])
    'EDFO 111 Group V/A&B' -> ('EDFO111', 'V', ['A', 'B'])  (group V = streams A and B)
    Returns None when the cell does not start with a unit code.
    """
    upper = " ".join(text.upper().split())
    m = _CODE_RE.match(upper)
    if not m:
        return None
    prefix, digits, rest = m.groups()
    group, subgroups = "", []
    g = _GROUP_RE.search(rest)
    if g:
        group = g.group(1)
        sub = re.match(r"\s*/\s*([A-Z](?:\s*&\s*[A-Z])*)\b", rest[g.end():])
        if sub:
            subgroups = re.findall(r"[A-Z]", sub.group(1))
    return f"{prefix}{digits}", group, subgroups


def _parse_group_definitions(text: str) -> dict[str, list[str]]:
    """'Group B: KISW/RELI, PE/RELI' -> {'B': ['KISW/RELI', 'PE/RELI']}"""
    groups: dict[str, list[str]] = {}
    for line in text.splitlines():
        m = _GROUP_DEF_RE.match(line.strip())
        if not m:
            continue
        letter, body = m.group(1).upper(), m.group(2)
        body = re.sub(r"\(.*?\)", " ", body)  # drop "(0745...-collins)" class-rep notes
        combos = re.findall(r"[A-Z]{2,6}\s*/\s*[A-Z]{2,6}(?:\s*/\s*[A-Z]{2,6})?", body.upper())
        if combos:
            groups[letter] = [re.sub(r"\s+", "", c) for c in combos]
    return groups


def clean_lecturer_name(raw: str) -> str:
    """'Dr. Joseph Omollo FT) 0781977790' -> 'Dr. Joseph Omollo'"""
    name = _PHONE_RE.sub(" ", raw)
    name = re.sub(r"\(?\b(?:FT|PT)\b\)?", " ", name, flags=re.IGNORECASE)
    name = re.sub(r"\(.*?\)", " ", name)
    name = re.sub(r"\s*/\s*", " / ", name)
    name = " ".join(name.replace("\n", " ").split())
    return name.strip(" /-,")


def is_placeholder_lecturer(name: str) -> bool:
    return not name or bool(_PLACEHOLDER_RE.match(name.strip()))


class _Context:
    def __init__(self):
        self.heading = ""
        self.programme = ""
        self.year: int | None = None
        self.semester: int | None = None
        self.odel = False
        self.odel_part = False
        self.groups: dict[str, list[str]] = {}

    def update_from_text(self, text: str, paragraph: bool = False) -> None:
        text = text.strip()
        if not text:
            return
        first_line = text.splitlines()[0].strip()
        part = _PART_RE.match(first_line)
        if part:
            self.odel_part = part.group(1).upper() == "ODEL"
            return
        m = _YEAR_SEM_RE.search(text)
        if (
            paragraph and not m and first_line.isupper() and len(first_line) > 8
            and (not _PROGRAMME_RE.match(first_line) or first_line.rstrip(" :").endswith("PROGRAMMES"))
            and not _GROUP_DEF_RE.match(first_line)
        ):
            # Part titles ("POSTGRADUATE PROGRAMMES", "SERVICING FOR OTHER DEPARTMENTS",
            # "DEPARTMENT OF EDUCATION SEP-DEC 2026 (ODEL)") start or end an ODEL part.
            self.odel_part = bool(_ODEL_RE.search(first_line))
        if _PROGRAMME_RE.match(first_line):
            # "BACHELOR OF SCIENCE IN ANIMAL SCIENCE Y1S1 SEP 2026 ..." -> programme title only
            ym = _YEAR_SEM_RE.search(first_line)
            self.programme = first_line[: ym.start()] if ym else first_line
            self.programme = self.programme.strip(" :,-–")[:200]
            if not m:
                # New programme whose year/semester follows on its own line.
                self.heading, self.year, self.semester, self.groups = first_line[:200], None, None, {}
                self.odel = self.odel_part or bool(_ODEL_RE.search(first_line))
        # A new year/semester heading starts a new section. Group definitions
        # usually sit on the lines right after it (same cell or next paragraphs).
        if m and not parse_code_cell(text):
            self.heading = first_line[:200]
            self.year, self.semester = int(m.group(1)), int(m.group(2))
            self.odel = self.odel_part or bool(_ODEL_RE.search(text) or _ODEL_RE.search(self.programme))
            self.groups = {}
        self.groups.update(_parse_group_definitions(text))


def _find_header(texts: list[str]) -> dict[str, int] | None:
    layout: dict[str, int] = {}
    for idx, text in enumerate(texts):
        t = text.lower()
        if "code" in t and "code" not in layout:
            layout["code"] = idx
        elif ("title" in t or t.strip() in {"unit", "units", "unit name", "course"}) and "title" not in layout:
            layout["title"] = idx
        elif ("lecturer" in t or "instructor" in t) and "lecturer" not in layout:
            layout["lecturer"] = idx
    if "code" in layout and "lecturer" in layout:
        return layout
    return None


def _iter_body(doc):
    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield "p", Paragraph(el, doc)
        elif tag == "tbl":
            yield "tbl", Table(el, doc)


def detect_department(file_path: str) -> str:
    """
    The department an allocation document comes from, read from its letterhead
    ('DEPARTMENT OF EDUCATION'), or '' when the document does not say.
    """
    for kind, block in _iter_body(Document(file_path)):
        if kind == "tbl":
            break
        m = _DEPARTMENT_RE.match(block.text.strip())
        if m:
            return " ".join(m.group(1).upper().split())
    return ""


def parse_allocation_docx(file_path: str) -> list[dict]:
    """
    Parses a departmental allocation Word document (.docx).

    Walks paragraphs and tables in document order so each row inherits the
    programme section (year/semester, group combinations) it belongs to.
    Tables split across pages without their own header row reuse the last
    seen column layout.
    """
    doc = Document(file_path)
    ctx = _Context()
    layout: dict[str, int] | None = None
    allocations: list[dict] = []

    for table_idx, (kind, block) in enumerate(_iter_body(doc)):
        if kind == "p":
            ctx.update_from_text(block.text, paragraph=True)
            continue

        for row_idx, row in enumerate(block.rows):
            cells = _unique_cells(row)
            texts = [c.text.strip() for c in cells]

            header = _find_header(texts)
            if header:
                layout = header
                continue

            code_idx = layout["code"] if layout else 0
            parsed = parse_code_cell(texts[code_idx]) if len(texts) > code_idx else None
            if not parsed:
                # Section heading / group definitions / "COMMON UNITS" rows.
                ctx.update_from_text("\n".join(dict.fromkeys(texts)))
                continue
            if not layout or len(texts) <= layout["lecturer"]:
                continue

            unit_code, group, subgroups = parsed
            title_idx = layout.get("title")
            raw_title = texts[title_idx] if title_idx is not None and len(texts) > title_idx else ""
            raw_lecturer = texts[layout["lecturer"]]
            # Rows whose merged cells shifted the lecturer one column right.
            if _PHONE_RE.fullmatch(raw_lecturer.replace(" ", "")) and layout["lecturer"] > 0:
                raw_lecturer = texts[layout["lecturer"] - 1]

            name = clean_lecturer_name(raw_lecturer)
            allocations.append({
                "unit_code": unit_code,
                "raw_unit_code": " ".join(texts[code_idx].split()),
                "unit_title": " ".join(raw_title.split()),
                "lecturer_name": name,
                "placeholder": is_placeholder_lecturer(name),
                "group": group,
                "subgroups": subgroups,
                "group_combinations": ctx.groups.get(group, []) if group else [],
                "year_of_study": ctx.year,
                "semester": ctx.semester,
                "odel": ctx.odel,
                "section": ctx.heading,
                "programme": ctx.programme,
                "source_file": os.path.basename(file_path),
                "source": f"{os.path.basename(file_path)} table-block {table_idx} row {row_idx}",
            })

    return allocations


def convert_doc_to_docx(file_path: str) -> str:
    """Converts a legacy .doc file with LibreOffice; returns the .docx path."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise ValueError("Legacy .doc files need LibreOffice on the server; please save the file as .docx.")
    out_dir = tempfile.mkdtemp(prefix="alloc_doc_")
    subprocess.run(
        [soffice, "--headless", "--convert-to", "docx", "--outdir", out_dir, file_path],
        check=True, capture_output=True, timeout=120,
    )
    converted = os.path.join(out_dir, os.path.splitext(os.path.basename(file_path))[0] + ".docx")
    if not os.path.exists(converted):
        raise ValueError("LibreOffice could not convert the .doc file.")
    return converted


def parse_allocation_pdf(file_obj) -> list[dict]:
    """PDF parser fallback if a PDF allocation file is uploaded."""
    import pdfplumber

    with pdfplumber.open(file_obj) as pdf:
        lines = [line for page in pdf.pages for line in (page.extract_text() or "").split("\n")]
    ctx = _Context()
    rows = []
    for line in lines:
        match = re.match(r"\s*([A-Z]{3,4}\s*\d{3,5}(?:\s*(?:GROUP|GRP|GR\.?|G)\s*[A-Z]\d?\b)?)\s+(.+?)\s+([A-Z][a-zA-Z\s\.\(\)\/’']+?)\s*(?:(?:\+?254|0)[17]\d{8})?\s*$", line)
        parsed = parse_code_cell(match.group(1)) if match else None
        if not parsed:
            ctx.update_from_text(line)
            continue
        unit_code, group, subgroups = parsed
        name = clean_lecturer_name(match.group(3))
        rows.append({
            "unit_code": unit_code,
            "raw_unit_code": match.group(1).strip(),
            "unit_title": match.group(2).strip(),
            "lecturer_name": name,
            "placeholder": is_placeholder_lecturer(name),
            "group": group,
            "subgroups": subgroups,
            "group_combinations": ctx.groups.get(group, []) if group else [],
            "year_of_study": ctx.year,
            "semester": ctx.semester,
            "odel": ctx.odel,
            "section": ctx.heading,
            "programme": ctx.programme,
            "source_file": "pdf",
            "source": "pdf",
        })
    return rows
