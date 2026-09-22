import os
import tempfile
from datetime import date, time

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from django.urls import reverse
from docx import Document
from rest_framework.test import APITestCase

from apps.accounts.models import User
from apps.departments.models import Department, Faculty
from apps.lecturers.models import Lecturer
from apps.programs.models import Program
from apps.rooms.models import Room
from apps.timetable.models import AcademicTerm, AllocationDocument, TimetableSlot
from apps.timetable.services.allocation_matcher import (
    SlotRef,
    display_lecturer,
    plan_assignments,
    program_matches_heading,
)
from apps.timetable.services.allocation_parser import (
    clean_lecturer_name,
    parse_allocation_docx,
    parse_code_cell,
)
from apps.units.models import Unit

HEADER = ["Course code", "Course Title", "L", "P", "CF", "LECTURER", "PHONE NO."]


def build_allocation_docx(sections, department: str = "") -> bytes:
    """sections: [(heading_lines, [row, ...]), ...]; each row is a 7-cell list."""
    doc = Document()
    if department:
        doc.add_paragraph(department)
    for heading_lines, rows in sections:
        for line in heading_lines:
            doc.add_paragraph(line)
        table = doc.add_table(rows=0, cols=len(HEADER))
        for values in [HEADER] + rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, values):
                cell.text = value
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        path = f.name
    try:
        doc.save(path)
        with open(path, "rb") as f:
            return f.read()
    finally:
        os.remove(path)


def parse_bytes(data: bytes) -> list[dict]:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        f.write(data)
        path = f.name
    try:
        return parse_allocation_docx(path)
    finally:
        os.remove(path)


BED_ARTS_Y1 = [
    "BACHELOR OF EDUCATION(ARTS): Y1S1 N=2000 INTAKE: SEP 2026",
    "Group A: ENGL/LIT",
    "Group B: KISW/RELI, PE/RELI",
    "Group L: ENGL/LIT",
]


class AllocationParserTests(SimpleTestCase):
    def test_code_cell_variants(self):
        self.assertEqual(parse_code_cell("COSC 103 Group A"), ("COSC103", "A", []))
        self.assertEqual(parse_code_cell("BOTA 101 G B"), ("BOTA101", "B", []))
        self.assertEqual(parse_code_cell("EDCI 211 GA"), ("EDCI211", "A", []))
        self.assertEqual(parse_code_cell("MATH 122 (GRP K)"), ("MATH122", "K", []))
        self.assertEqual(parse_code_cell("EDFO 111Group W/C"), ("EDFO111", "W", ["C"]))
        self.assertEqual(parse_code_cell("EDFO 111\nGroup V/A&B"), ("EDFO111", "V", ["A", "B"]))
        self.assertEqual(parse_code_cell("COSC 00104"), ("COSC00104", "", []))
        self.assertEqual(parse_code_cell("COSC 103 Group D2"), ("COSC103", "D2", []))
        self.assertIsNone(parse_code_cell("Course code"))

    def test_lecturer_name_cleaning(self):
        self.assertEqual(clean_lecturer_name("Dr. Joseph Omollo FT)"), "Dr. Joseph Omollo")
        self.assertEqual(clean_lecturer_name("Luke Mwema/ Kevin Tuei (FT) 0798382994"), "Luke Mwema / Kevin Tuei")

    def test_rows_carry_section_context_and_placeholders(self):
        data = build_allocation_docx([
            (BED_ARTS_Y1, [
                ["COSC 103 Group A", "ICT Skills", "30", "15", "3.0", "Joseph Mutwiri (FT)", "0701656589"],
                ["COSC 103 Group B", "ICT Skills", "30", "15", "3.0", "Michael Mutisya (FT)", "0720752707"],
                ["ETHI 101", "Ethics", "15", "30", "3.0", "DHUM- 1.2", ""],
            ]),
        ])
        rows = parse_bytes(data)
        self.assertEqual(len(rows), 3)
        a, b, ethi = rows
        self.assertEqual((a["year_of_study"], a["semester"]), (1, 1))
        self.assertEqual(a["programme"], "BACHELOR OF EDUCATION(ARTS)")
        self.assertEqual(a["group_combinations"], ["ENGL/LIT"])
        self.assertEqual(b["group_combinations"], ["KISW/RELI", "PE/RELI"])
        self.assertTrue(ethi["placeholder"])
        self.assertFalse(a["placeholder"])

    def test_headerless_continuation_table_reuses_layout(self):
        doc = Document()
        doc.add_paragraph("BACHELOR OF SCIENCE IN AGRICULTURE Y3S1 2024 -SEP INTAKE")
        header = doc.add_table(rows=1, cols=7)
        for cell, value in zip(header.rows[0].cells, HEADER):
            cell.text = value
        body = doc.add_table(rows=1, cols=7)
        for cell, value in zip(body.rows[0].cells, ["AGEN 352", "Irrigation", "30", "30", "3", "Fredrick Marega (PT)", "0702122152"]):
            cell.text = value
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            path = f.name
        try:
            doc.save(path)
            rows = parse_allocation_docx(path)
        finally:
            os.remove(path)
        self.assertEqual([(r["unit_code"], r["lecturer_name"], r["year_of_study"]) for r in rows],
                         [("AGEN352", "Fredrick Marega", 3)])


class AllocationMatcherTests(SimpleTestCase):
    def test_programme_heading_matching(self):
        self.assertTrue(program_matches_heading("BSC.ANIMAL SCI", "BACHELOR OF SCIENCE IN ANIMAL SCIENCE"))
        self.assertTrue(program_matches_heading("BSC.AGED", "BACHELOR OF SCIENCE IN AGRICULTURAL EDUCATION AND EXTENSION"))
        self.assertTrue(program_matches_heading("DIP.T&H MNGT", "DIPLOMA IN TOURISM AND HOTEL MANAGEMENT"))
        self.assertTrue(program_matches_heading("BED.KISW/RELI", "BACHELOR OF EDUCATION(ARTS)"))
        self.assertFalse(program_matches_heading("DIP.IRM", "DIPLOMA IN INFORMATION TECHNOLOGY"))
        self.assertFalse(program_matches_heading("BSC MATH/COMP", "BACHELOR OF SCIENCE IN BIOLOGY"))

    def _row(self, code, lecturer, **extra):
        row = {"unit_code": code, "raw_unit_code": code, "lecturer_name": lecturer, "placeholder": False,
               "group": "", "subgroups": [], "group_combinations": [], "year_of_study": 1, "semester": 1,
               "odel": False, "section": "", "programme": "", "source_file": "a.docx"}
        row.update(extra)
        return row

    def test_no_substring_code_matches(self):
        slots = [SlotRef(1, "COSC104", "BSC COMP SCI", 1, 1), SlotRef(2, "COSC1040", "BSC COMP SCI", 1, 1)]
        plan = plan_assignments([self._row("COSC00104", "Luke Mwema")], slots)
        self.assertEqual(set(plan["assignments"]), {1})

    def test_group_combinations_pin_programmes(self):
        slots = [
            SlotRef(1, "COSC103", "BED.ENG/LIT", 1, 1),
            SlotRef(2, "COSC103", "BED.KISW/RELI", 1, 1),
            SlotRef(3, "COSC103", "BED.PE/RELI", 1, 1),
            SlotRef(4, "COSC103", "BSC COMP SCI", 1, 1),
        ]
        rows = [
            self._row("COSC103", "Joseph Mutwiri", group="A", group_combinations=["ENGL/LIT"],
                      programme="BACHELOR OF EDUCATION(ARTS)"),
            self._row("COSC103", "Michael Mutisya", group="B", group_combinations=["KISW/RELI", "PE/RELI"],
                      programme="BACHELOR OF EDUCATION(ARTS)"),
        ]
        plan = plan_assignments(rows, slots)
        got = {sid: r["lecturer_name"] for sid, (r, _) in plan["assignments"].items()}
        self.assertEqual(got, {1: "Joseph Mutwiri", 2: "Michael Mutisya", 3: "Michael Mutisya"})

    def test_year_separates_same_unit(self):
        slots = [SlotRef(1, "COSC221", "BSC COMP SCI", 2, 1), SlotRef(2, "COSC221", "BSC.INFORM SCI", 4, 1)]
        rows = [
            self._row("COSC221", "Kenneth Gitonga", year_of_study=2, programme="BACHELOR OF SCIENCE IN COMPUTER SCIENCE"),
            self._row("COSC221", "Michael Mutisya", year_of_study=4, programme="BACHELOR OF SCIENCE IN INFORMATION SCIENCE"),
        ]
        got = {sid: r["lecturer_name"] for sid, (r, _) in plan_assignments(rows, slots)["assignments"].items()}
        self.assertEqual(got, {1: "Kenneth Gitonga", 2: "Michael Mutisya"})

    def test_groups_sharing_a_programme_map_to_streams(self):
        slots = [SlotRef(1, "COSC103", "BED.ENG/LIT", 1, 1, stream="1"), SlotRef(2, "COSC103", "BED.ENG/LIT", 1, 1, stream="2")]
        rows = [
            self._row("COSC103", "Joseph Mutwiri", group="A", group_combinations=["ENGL/LIT"], programme="BACHELOR OF EDUCATION(ARTS)"),
            self._row("COSC103", "Antony Mwangi", group="L", group_combinations=["ENGL/LIT"], programme="BACHELOR OF EDUCATION(ARTS)"),
        ]
        got = {sid: r["lecturer_name"] for sid, (r, _) in plan_assignments(rows, slots)["assignments"].items()}
        self.assertEqual(got, {1: "Joseph Mutwiri", 2: "Antony Mwangi"})

    def test_unresolvable_disagreement_is_reported_not_overwritten(self):
        slots = [SlotRef(1, "EDCI211", "BED.ECDE", 2, 1)]
        rows = [
            self._row("EDCI211", "Dr. Obote Denis", year_of_study=2, programme="BACHELOR OF EDUCATION(ECDE)"),
            self._row("EDCI211", "Dr. Rose Mugwiria", year_of_study=2, programme="BACHELOR OF EDUCATION(ECDE)"),
        ]
        plan = plan_assignments(rows, slots)
        self.assertEqual(plan["assignments"], {})
        self.assertEqual(len(plan["conflicts"]), 1)

    def test_unit_not_scheduled_for_programme_does_not_spill(self):
        slots = [SlotRef(1, "COSC103", "BED.ENG/LIT", 1, 1), SlotRef(2, "PHIL104", "BSC.AGRICULTURE", 1, 1)]
        plan = plan_assignments([self._row("COSC103", "Joseph Mutwiri", programme="BACHELOR OF SCIENCE IN AGRICULTURE")], slots)
        self.assertEqual(plan["assignments"], {})
        self.assertIn("not scheduled", plan["skipped"][0][1])

    def test_co_taught_display_keeps_both_names(self):
        self.assertEqual(display_lecturer("Kevin Tuei", "Luke Mwema / Kevin Tuei"), "Luke Mwema / Kevin Tuei")
        self.assertEqual(display_lecturer("Old Wrong", "Joseph Mutwiri"), "Old Wrong")
        self.assertEqual(display_lecturer("", "Joseph Mutwiri"), "Joseph Mutwiri")

    def test_placeholder_and_odel_rows_skipped(self):
        slots = [SlotRef(1, "ETHI101", "BSC COMP SCI", 1, 1)]
        rows = [self._row("ETHI101", "DHUM", placeholder=True), self._row("ETHI101", "Someone", odel=True)]
        self.assertEqual(plan_assignments(rows, slots)["assignments"], {})


COSC_ROW_A = ["COSC 103 Group A", "ICT Skills", "30", "15", "3.0", "Joseph Mutwiri (FT)", "0701656589"]
COSC_ROW_B = ["COSC 103 Group B", "ICT Skills", "30", "15", "3.0", "Michael Mutisya (FT)", "0720752707"]
DCOMP = "DEPARTMENT OF COMPUTER SCIENCE AND ICT"
DEDU = "DEPARTMENT OF EDUCATION"


class AssignLecturersAPITests(APITestCase):
    def setUp(self):
        faculty = Faculty.objects.create(name="Education", code="EDU")
        self.dept = Department.objects.create(faculty=faculty, name="Education", code="DEDU")
        self.term = AcademicTerm.objects.create(
            academic_year="2026/2027", semester=1, start_date=date(2026, 9, 1),
            end_date=date(2026, 12, 20), is_current=True,
        )
        self.room = Room.objects.create(code="UTC 1", name="UTC 1", building="UTC", capacity=100)
        self.unit = Unit.objects.create(code="COSC103", name="Introduction to ICT Skills", credit_hours=3, department=self.dept)
        self.eng = Program.objects.create(code="BED.ENG/LIT", name="BED.ENG/LIT", department=self.dept)
        self.kisw = Program.objects.create(code="BED.KISW/RELI", name="BED.KISW/RELI", department=self.dept)
        self.geog = Program.objects.create(code="BED.GEOG/HIST", name="BED.GEOG/HIST", department=self.dept)

        self.admin = User.objects.create_user(username="admin", password="x", role=User.Role.ADMIN, university_id="ADM1", is_superuser=True)
        mutwiri = User.objects.create_user(username="jm", password="x", role=User.Role.LECTURER, university_id="L1",
                                           first_name="Joseph", last_name="Mutwiri")
        self.mutwiri = Lecturer.objects.create(user=mutwiri, department=self.dept)
        other = User.objects.create_user(username="ow", password="x", role=User.Role.LECTURER, university_id="L2",
                                         first_name="Old", last_name="Wrong")
        self.wrong = Lecturer.objects.create(user=other, department=self.dept)

        def slot(program, day, lecturer=None):
            return TimetableSlot.objects.create(
                term=self.term, unit=self.unit, program=program, year_of_study=1, room=self.room,
                day_of_week=day, start_time=time(7), end_time=time(9), lecturer=lecturer,
            )
        self.eng_slot = slot(self.eng, "mon")
        self.kisw_slot = slot(self.kisw, "tue", lecturer=self.wrong)
        self.geog_slot = slot(self.geog, "wed")
        self.client.force_authenticate(self.admin)

    def _post(self, rows, department=DCOMP, name="alloc.docx", **extra):
        data = build_allocation_docx([(BED_ARTS_Y1, rows)], department=department)
        upload = SimpleUploadedFile(name, data)
        res = self.client.post(reverse("timetable-assign-lecturers"), {"file": upload, **extra}, format="multipart")
        self.assertEqual(res.status_code, 200, getattr(res, "data", res))
        return res.data

    def _names(self):
        return {s.program.code: s.lecturer_name_text for s in TimetableSlot.objects.select_related("program")}

    def test_dry_run_reports_without_saving(self):
        data = self._post([COSC_ROW_A, COSC_ROW_B], dry_run="true")
        self.assertEqual(data["slots_assigned_by_this_document"], 2)
        self.assertEqual(data["slots_updated"], 0)
        self.assertFalse(AllocationDocument.objects.exists())
        self.eng_slot.refresh_from_db()
        self.assertEqual(self.eng_slot.lecturer_name_text, "")

    def test_assigns_exact_slots_and_clears_stale_account(self):
        data = self._post([COSC_ROW_A, COSC_ROW_B])
        self.assertEqual(data["source"], DCOMP)
        self.eng_slot.refresh_from_db()
        self.kisw_slot.refresh_from_db()
        self.assertEqual(self.eng_slot.lecturer_name_text, "Joseph Mutwiri")
        self.assertEqual(self.eng_slot.lecturer_id, self.mutwiri.id)
        self.assertEqual(self.kisw_slot.lecturer_name_text, "Michael Mutisya")
        self.assertIsNone(self.kisw_slot.lecturer_id)
        self.assertEqual(self.geog_slot.lecturer_name_text, "")

    def test_later_department_does_not_overwrite_more_specific_rows(self):
        self._post([COSC_ROW_A, COSC_ROW_B], department=DCOMP)
        # Another department lists COSC 103 for the whole programme without a group:
        # it may fill the class nobody else covers, but not the group-specific ones.
        general = ["COSC 103", "ICT Skills", "30", "15", "3.0", "Ann Other (PT)", ""]
        data = self._post([general], department=DEDU, name="dedu.docx")
        self.assertEqual(self._names(), {
            "BED.ENG/LIT": "Joseph Mutwiri", "BED.KISW/RELI": "Michael Mutisya", "BED.GEOG/HIST": "Ann Other",
        })
        self.assertEqual(data["departments_on_file"], [DCOMP, DEDU])

    def test_equal_clash_between_departments_is_reported_and_slot_kept(self):
        self._post([COSC_ROW_A], department=DCOMP)
        clash = ["COSC 103 Group A", "ICT Skills", "30", "15", "3.0", "Somebody Else (PT)", ""]
        data = self._post([clash], department=DEDU, name="dedu.docx")
        self.assertEqual(data["conflict_slots"], 1)
        self.assertEqual({c["department"] for c in data["conflicts"][0]["candidates"]}, {DCOMP, DEDU})
        self.eng_slot.refresh_from_db()
        self.assertEqual(self.eng_slot.lecturer_name_text, "Joseph Mutwiri")

    def test_reupload_replaces_that_departments_previous_version(self):
        self._post([COSC_ROW_A, COSC_ROW_B], department=DCOMP, name="v1.docx")
        data = self._post([COSC_ROW_A], department=DCOMP, name="v2.docx")
        self.assertIn("replacing", data["detail"])
        self.assertEqual(AllocationDocument.objects.count(), 1)
        self.assertEqual(self._names()["BED.KISW/RELI"], "")  # row dropped in v2
        self.assertEqual(self._names()["BED.ENG/LIT"], "Joseph Mutwiri")

    def test_delete_department_clears_its_slots(self):
        self._post([COSC_ROW_A], department=DCOMP)
        self._post([COSC_ROW_B], department=DEDU, name="dedu.docx")
        listing = self.client.get(reverse("timetable-assign-lecturers")).data
        self.assertEqual({d["source"]: d["slots_assigned"] for d in listing["documents"]}, {DCOMP: 1, DEDU: 1})
        res = self.client.delete(reverse("timetable-assign-lecturers") + f"?source={DCOMP}")
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(self._names()["BED.ENG/LIT"], "")
        self.assertEqual(self._names()["BED.KISW/RELI"], "Michael Mutisya")
